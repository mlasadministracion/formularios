#!/usr/bin/env python3
"""
Uso: python3 generar_recibos.py <nombre_hoja>
Genera un recibo PDF por cada estudiante en la hoja indicada.
"""

import sys
import os
import re
import io
import smtplib
from email.message import EmailMessage
import openpyxl
from openpyxl.utils import get_column_letter
from weasyprint import HTML
import requests
import config

ARCHIVO_EXCEL = "https://docs.google.com/spreadsheets/d/1w7X_XNzXVqcvLIP5UThRDaPcZPwzbTXa/export?format=xlsx"
TEMPLATE_HTML = "recibo_template.html"
CARPETA_BASE = "recibos"
COLUMNA_FIN = "V"
FILA_ETIQUETAS = 3
FILA_BUSQUEDA = 4

CAMPOS_MONETARIOS = {"pension", "almuerzo", "pago 1", "pago 2", "pago 3"}

def col_letra_a_num(letra):
    """Convierte una letra de columna de Excel (ej. 'V') a su número (22)."""
    num = 0
    for c in letra.upper():
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


def formatear_valor(val, etiqueta=""):
    """Da formato de presentación a un valor de celda: '-' si está vacío/NA,
    separador de miles con puntos, y prefijo '$' si la etiqueta es un campo monetario."""
    if val is None or str(val).strip().upper() == "NA":
        return "-"
    if isinstance(val, float) and val == int(val):
        val = int(val)
    if isinstance(val, (int, float)):
        numero = f"{int(val):,}".replace(",", ".")
        if etiqueta in CAMPOS_MONETARIOS:
            return f"${numero}"
        return numero
    return str(val).strip()


def cargar_workbook():
    """Descarga el Excel publicado en Google Sheets (export?format=xlsx) y lo
    carga en memoria con openpyxl, sin escribir nada a disco."""
    print("Descargando hoja de cálculo desde Google Sheets...")
    resp = requests.get(ARCHIVO_EXCEL, timeout=30)
    resp.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)


def leer_hoja(nombre_hoja):
    """Lee la hoja del mes indicado y devuelve la lista de estudiantes a facturar.

    - Las etiquetas de cada columna se toman de la fila 3 (FILA_ETIQUETAS).
    - La primera fila con un número en la columna A marca el inicio de los datos.
    - Se ignoran filas sin nombre de estudiante (columna B), como las de totales.
    - Se ignoran estudiantes cuya columna "genera factura" sea distinta de "si"
      (si la columna no existe en la hoja, se asume "si" para todos).
    """
    wb = cargar_workbook()
    if nombre_hoja not in wb.sheetnames:
        print(f"Error: la hoja '{nombre_hoja}' no existe.")
        print(f"Hojas disponibles: {', '.join(wb.sheetnames)}")
        sys.exit(1)
    ws = wb[nombre_hoja]
    max_col = col_letra_a_num(COLUMNA_FIN)

    etiquetas = {}
    for col in range(1, max_col + 1):
        val = ws.cell(row=FILA_ETIQUETAS, column=col).value
        if val:
            etiquetas[get_column_letter(col)] = str(val).strip().lower()

    fila_inicio = None
    for row in ws.iter_rows(min_row=FILA_BUSQUEDA, max_col=1):
        if isinstance(row[0].value, (int, float)):
            fila_inicio = row[0].row
            break

    if fila_inicio is None:
        print("No se encontraron estudiantes.")
        sys.exit(1)

    estudiantes = []
    for fila in ws.iter_rows(min_row=fila_inicio, max_col=max_col):
        num = fila[0].value
        if num is None or not isinstance(num, (int, float)):
            break
        if not fila[1].value:
            continue
        registro = {}
        for cell in fila:
            letra = get_column_letter(cell.column)
            if letra in etiquetas:
                registro[etiquetas[letra]] = cell.value
        genera = str(registro.get("genera factura", "si")).strip().lower()
        if genera != "si":
            continue
        estudiantes.append(registro)
    return estudiantes


def generar_pdf(estudiante, carpeta_salida, indice, template_html, nombre_hoja):
    """Rellena la plantilla HTML con los datos del estudiante y la convierte a PDF.

    Reemplaza cada placeholder {etiqueta} por su valor formateado, además de
    {mes_upper}/{mes_title} por el nombre de la hoja. Si el almuerzo es "-" o
    "$0" se elimina la fila de almuerzo de la tabla de conceptos. El PDF se
    guarda en carpeta_salida con el nombre del estudiante y se devuelve su ruta.
    """
    reemplazos = {}
    for etiqueta, valor in estudiante.items():
        reemplazos[etiqueta] = formatear_valor(valor, etiqueta)

    html = template_html
    for clave, valor in reemplazos.items():
        html = html.replace("{" + clave + "}", valor)
    html = html.replace("{mes_upper}", nombre_hoja.upper())
    html = html.replace("{mes_title}", nombre_hoja.capitalize())

    almuerzo_val = reemplazos.get("almuerzo", "-")
    if almuerzo_val in ("-", "$0"):
        html = re.sub(r'<tr data-id="almuerzo".*?</tr>', '', html, flags=re.DOTALL)

    nombre_estudiante = estudiante.get("estudiante", f"estudiante_{indice}")
    nombre_estudiante = re.sub(r'[\\/*?:"<>|]', "", str(nombre_estudiante)).strip()
    archivo_pdf = os.path.join(carpeta_salida, f"{nombre_estudiante}.pdf")

    base_url = os.path.dirname(os.path.abspath(__file__))
    HTML(string=html, base_url=base_url).write_pdf(archivo_pdf)
    return archivo_pdf


def enviar_email(archivo_pdf, destinatarios, mes, estudiante):
    """Envía el PDF del recibo como adjunto por correo (Gmail SMTP_SSL) a la
    lista de destinatarios. Asunto y cuerpo se toman de config.py. Devuelve
    False sin hacer nada si no hay destinatarios."""
    if not destinatarios:
        return False

    msg = EmailMessage()
    msg["From"] = f"{config.EMAIL_NOMBRE} <{config.EMAIL_REMITENTE}>"
    msg["To"] = ", ".join(destinatarios)
    msg["Subject"] = config.EMAIL_ASUNTO.format(mes=mes.upper(), estudiante=estudiante.upper())
    msg.set_content(config.EMAIL_CUERPO.format(mes=mes))

    with open(archivo_pdf, "rb") as f:
        msg.add_attachment(f.read(), maintype="application", subtype="pdf",
                           filename=os.path.basename(archivo_pdf))

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
        smtp.login(config.EMAIL_REMITENTE, config.EMAIL_PASSWORD)
        smtp.send_message(msg)

    return True


def main():
    """Punto de entrada CLI: lee argumentos, genera un PDF por estudiante y,
    si se pasó --enviar, los envía por correo a las direcciones de la columna
    "correo". Uso: python3 generar_recibos.py <nombre_hoja> [--enviar]"""
    if len(sys.argv) < 2:
        print("Uso: python3 generar_recibos.py <nombre_hoja> [--enviar]")
        sys.exit(1)

    enviar = "--enviar" in sys.argv
    args = [a for a in sys.argv[1:] if a != "--enviar"]
    nombre_hoja = " ".join(args)
    estudiantes = leer_hoja(nombre_hoja)

    with open(TEMPLATE_HTML, "r", encoding="utf-8") as f:
        template_html = f.read()

    carpeta_salida = os.path.join(CARPETA_BASE, nombre_hoja)
    os.makedirs(carpeta_salida, exist_ok=True)

    print(f"\nGenerando recibos para hoja '{nombre_hoja}'...\n")
    enviados = 0
    for i, est in enumerate(estudiantes, 1):
        nombre = (est.get("estudiante") or f"Estudiante {i}").strip()
        archivo = generar_pdf(est, carpeta_salida, i, template_html, nombre_hoja)

        correo_raw = est.get("correo", "")
        destinatarios = [e.strip() for e in str(correo_raw).split(",") if "@" in e.strip()]

        print(f"  ✓ {nombre} → {archivo}")

        if enviar:
            if destinatarios:
                try:
                    enviar_email(archivo, destinatarios, nombre_hoja, nombre)
                    print(f"    Enviado a: {', '.join(destinatarios)}")
                    enviados += 1
                except Exception as e:
                    print(f"    Error al enviar: {e}")
            else:
                print(f"    Sin correo registrado")

    resumen = f"{len(estudiantes)} recibo(s) generado(s)"
    if enviar:
        resumen += f", {enviados} enviado(s) por correo"
    print(f"\nTotal: {resumen} — carpeta '{carpeta_salida}/'")


if __name__ == "__main__":
    main()
